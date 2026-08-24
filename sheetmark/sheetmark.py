#!/usr/bin/env python3
"""Convert local Excel workbooks to Markdown tables."""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable, Sequence
from urllib.parse import quote


SUPPORTED_DIRECT = {".xlsx", ".xlsm"}
SUPPORTED_LEGACY = {".xls"}


def find_soffice() -> str | None:
    """Find LibreOffice without assuming a platform-specific PATH entry."""
    candidates: list[str] = []
    configured = os.environ.get("SHEETMARK_SOFFICE")
    if configured:
        candidates.append(configured)

    for command in ("soffice", "libreoffice"):
        found = shutil.which(command)
        if found:
            candidates.append(found)

    if sys.platform == "darwin":
        candidates.extend(
            [
                "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                str(Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"),
            ]
        )
    elif sys.platform == "win32":
        for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
            root = os.environ.get(variable)
            if root:
                candidates.append(str(Path(root) / "LibreOffice" / "program" / "soffice.exe"))
                candidates.append(str(Path(root) / "LibreOffice" / "program" / "soffice.com"))

    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
    return None


def require_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Required local Python package is missing: openpyxl.\n"
            "Install it, then run SheetMark again:\n"
            "  macOS/Linux: python3 -m pip install openpyxl\n"
            "  Windows:     py -m pip install openpyxl"
        ) from error
    return load_workbook


def convert_legacy_xls(input_path: Path, temporary_directory: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            "Required local application is missing: LibreOffice.\n"
            "The .xls format needs LibreOffice for local conversion. Install it, "
            "then run SheetMark again.\n"
            "macOS: install LibreOffice, or set SHEETMARK_SOFFICE to its soffice path.\n"
            "Windows: install LibreOffice, or set SHEETMARK_SOFFICE to soffice.exe."
        )

    output_directory = temporary_directory / "converted"
    profile_directory = temporary_directory / "profile"
    output_directory.mkdir()
    profile_directory.mkdir()
    profile_uri = profile_directory.as_uri()
    command = [
        soffice,
        f"-env:UserInstallation={profile_uri}",
        "--headless",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(output_directory),
        str(input_path),
    ]
    result = subprocess.run(command, capture_output=True, text=True)
    converted_path = output_directory / f"{input_path.stem}.xlsx"
    if result.returncode != 0 or not converted_path.is_file():
        details = (result.stderr or result.stdout).strip()
        raise RuntimeError(f"LibreOffice could not convert {input_path.name}. {details}")
    return converted_path


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(sep=" ")
    elif isinstance(value, (date, time)):
        text = value.isoformat()
    elif isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    else:
        text = str(value)
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")


def meaningful_bounds(rows: Sequence[Sequence[object]]) -> tuple[int, int] | None:
    last_row = 0
    last_column = 0
    for row_number, row in enumerate(rows, start=1):
        for column_number, value in enumerate(row, start=1):
            if value is not None and value != "":
                last_row = max(last_row, row_number)
                last_column = max(last_column, column_number)
    if not last_row or not last_column:
        return None
    return last_row, last_column


def worksheet_rows(worksheet) -> list[list[object]]:
    """Return values while keeping merged-cell content only at its top-left cell."""
    merged_cells = {
        (row, column)
        for merged_range in worksheet.merged_cells.ranges
        for row in range(merged_range.min_row, merged_range.max_row + 1)
        for column in range(merged_range.min_col, merged_range.max_col + 1)
        if (row, column) != (merged_range.min_row, merged_range.min_col)
    }
    return [
        [None if (cell.row, cell.column) in merged_cells else cell.value for cell in row]
        for row in worksheet.iter_rows()
    ]


def safe_filename(value: str) -> str:
    normalized = re.sub(r"[^\w.-]+", "-", value, flags=re.UNICODE).strip(".-")
    return normalized or "worksheet"


def export_images(worksheet, image_directory: Path) -> list[Path]:
    if not worksheet._images:
        return []

    image_directory.mkdir(parents=True, exist_ok=True)
    exported: list[Path] = []
    for index, image in enumerate(worksheet._images, start=1):
        image_format = (getattr(image, "format", None) or "png").lower()
        if image_format == "jpeg":
            image_format = "jpg"
        image_path = image_directory / f"{safe_filename(worksheet.title)}-{index}.{image_format}"
        image_path.write_bytes(image._data())
        exported.append(image_path)
    return exported


def markdown_destination(path: str) -> str:
    """Encode a relative path for use as a Markdown URL destination."""
    return quote(path, safe="/")


def worksheet_markdown(title: str, rows: Iterable[Sequence[object]], image_paths: Sequence[Path], output_path: Path) -> str:
    materialized = [list(row) for row in rows]
    bounds = meaningful_bounds(materialized)
    if not bounds:
        lines = [f"## {cell_text(title)}", "", "_(empty worksheet)_"]
    else:
        last_row, last_column = bounds
        trimmed = [row[:last_column] + [None] * max(0, last_column - len(row)) for row in materialized[:last_row]]
        header_index = next((index for index, row in enumerate(trimmed) if any(value not in (None, "") for value in row)), 0)
        header = trimmed[header_index]
        body = trimmed[header_index + 1 :]
        lines = [f"## {cell_text(title)}", "", "| " + " | ".join(cell_text(value) for value in header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
        lines.extend("| " + " | ".join(cell_text(value) for value in row) + " |" for row in body)

    for image_path in image_paths:
        relative_path = image_path.relative_to(output_path.parent).as_posix()
        lines.extend(["", f"![{cell_text(title)} image]({markdown_destination(relative_path)})"])
    return "\n".join(lines) + "\n"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert a local Excel workbook to Markdown.")
    parser.add_argument("input", type=Path, help="Path to a local .xlsx, .xlsm, or .xls file")
    parser.add_argument("-o", "--output", type=Path, help="Output Markdown file or existing directory")
    parser.add_argument("--sheet", action="append", help="Worksheet name; repeat for multiple worksheets")
    parser.add_argument("--values", action="store_true", help="Use cached formula results instead of formula text")
    parser.add_argument("--no-images", action="store_true", help="Do not export embedded worksheet images")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    if not input_path.is_file():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        return 2

    extension = input_path.suffix.lower()
    if extension not in SUPPORTED_DIRECT | SUPPORTED_LEGACY:
        print("Supported formats are .xlsx, .xlsm, and .xls.", file=sys.stderr)
        return 2

    output_path = args.output.expanduser() if args.output else input_path.with_suffix(".md")
    if output_path.exists() and output_path.is_dir():
        output_path = output_path / f"{input_path.stem}.md"
    image_directory = output_path.parent / f"{output_path.stem}_images"

    try:
        load_workbook = require_openpyxl()
        with tempfile.TemporaryDirectory(prefix="sheetmark-") as temporary_name:
            workbook_path = input_path
            if extension in SUPPORTED_LEGACY:
                workbook_path = convert_legacy_xls(input_path, Path(temporary_name))
            workbook = load_workbook(
                workbook_path,
                data_only=args.values,
                keep_vba=extension == ".xlsm",
            )
            requested_sheets = args.sheet or workbook.sheetnames
            missing = [name for name in requested_sheets if name not in workbook.sheetnames]
            if missing:
                raise RuntimeError(f"Worksheet not found: {', '.join(missing)}")
            sections = []
            for name in requested_sheets:
                worksheet = workbook[name]
                image_paths = [] if args.no_images else export_images(worksheet, image_directory)
                sections.append(worksheet_markdown(name, worksheet_rows(worksheet), image_paths, output_path))
            workbook.close()
    except RuntimeError as error:
        print(f"SheetMark could not continue:\n{error}", file=sys.stderr)
        return 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(sections), encoding="utf-8")
    print(f"Wrote Markdown: {output_path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
